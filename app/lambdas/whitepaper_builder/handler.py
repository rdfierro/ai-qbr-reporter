import io
import os
from datetime import datetime, timezone
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from ..common.logging import get_logger
from ..common.models import ArtifactRef, KpiPayload, NarrativeResponse, NarrativeSlide
from ..common.utils import (
    build_artifact_key,
    read_bytes_from_s3,
    resolve_execution_id,
    write_bytes_to_s3,
)

LOG = get_logger(__name__)

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def lambda_handler(event: Dict[str, Any], context: Any) -> Dict[str, Any]:
    kpi_data = event.get('kpis')
    narrative_data = event.get('narrative')
    charts = event.get('charts', [])

    if not kpi_data or not narrative_data:
        raise ValueError('KPI and narrative data are required for whitepaper builder')

    kpis = KpiPayload(**kpi_data)
    narrative = NarrativeResponse(**narrative_data)
    execution_id = resolve_execution_id(event, context)

    document = Document()

    _add_title_page(document, kpis)
    _add_metrics_overview(document, kpis)
    _add_charts(document, charts, kpis)
    _add_narrative_sections(document, narrative.slides)
    _add_citations(document, narrative.citations)
    _add_generated_footer(document, kpis)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    key = build_artifact_key('artifacts/whitepapers', execution_id, 'qbr_whitepaper.docx')
    write_bytes_to_s3(
        os.environ['ARTIFACTS_BUCKET'],
        key,
        buffer.read(),
        DOCX_CONTENT_TYPE,
    )

    LOG.info(
        'Whitepaper generated',
        extra={
            'extra_fields': {
                'executionId': execution_id,
                'sections': len(narrative.slides) + 3,
            }
        },
    )

    return {
        'kpis': kpi_data,
        'charts': charts,
        'narrative': narrative.dict(),
        'whitepaper': ArtifactRef(
            bucket=os.environ['ARTIFACTS_BUCKET'],
            key=key,
            content_type=DOCX_CONTENT_TYPE,
        ).dict(),
        'meta': {
            'executionId': execution_id,
            'generated_at': datetime.now(timezone.utc).isoformat(),
        },
    }


def _add_title_page(document: Document, kpis: KpiPayload) -> None:
    title = document.add_heading('Quarterly Business Review', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f'Period: {kpis.period}\n')
    run.bold = True
    subtitle.add_run(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d')}")

    document.add_page_break()


def _add_metrics_overview(document: Document, kpis: KpiPayload) -> None:
    document.add_heading('Key Metrics Overview', level=1)

    table = document.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Metric'
    header_cells[1].text = 'Value'

    rows = [
        ('Pipeline Coverage', f"{kpis.coverage.pipeline:.2f} / Target {kpis.coverage.target:.2f}"),
        (
            'Conversion Rates',
            f"MQL->SQL {kpis.conversion.mql_to_sql:.2%}; SQL->Win {kpis.conversion.sql_to_win:.2%}",
        ),
        (
            'Revenue',
            f"ARR ${kpis.revenue.arr:,}; New ${kpis.revenue.new:,}; Churn ${kpis.revenue.churn:,}",
        ),
        ('Customer Success', f"CSAT {kpis.support.csat:.2f}; SLA breaches {kpis.support.sla_breaches}"),
        (
            'Adoption',
            f"WAU {kpis.adoption.wau:,}; Feature X usage {kpis.adoption.feature_x_usage_pct:.2%}",
        ),
    ]

    for metric, value in rows:
        cells = table.add_row().cells
        cells[0].text = metric
        cells[1].text = value

    document.add_paragraph()


def _add_charts(document: Document, charts: List[Dict[str, Any]], kpis: KpiPayload) -> None:
    if not charts:
        return

    document.add_heading(f'Visualizations – {kpis.period}', level=1)
    bucket = os.environ['ARTIFACTS_BUCKET']

    for chart in charts:
        key = chart.get('key')
        if not key:
            continue
        image_bytes = read_bytes_from_s3(bucket, key)
        document.add_picture(io.BytesIO(image_bytes), width=Inches(6))
        caption = chart.get('title') or key.split('/')[-1]
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()


def _add_narrative_sections(document: Document, slides: List[NarrativeSlide]) -> None:
    for slide in slides:
        document.add_heading(slide.title, level=1)
        for bullet in slide.bullets:
            document.add_paragraph(bullet, style='List Bullet')
        document.add_paragraph()


def _add_citations(document: Document, citations: List[Dict[str, Any]]) -> None:
    if not citations:
        return

    document.add_heading('References', level=1)
    for citation in citations:
        text = citation.get('text') or citation.get('url') or str(citation)
        document.add_paragraph(text, style='List Number')

    document.add_paragraph()


def _add_generated_footer(document: Document, kpis: KpiPayload) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | Period {kpis.period}"
    ).italic = True

